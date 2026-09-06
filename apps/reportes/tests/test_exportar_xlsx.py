"""
La exportación en cada servicio, y en Excel con la línea gráfica de la UNL.

Dos cosas se pidieron sobre lo ya construido: que la exportación esté en todos
los servicios, y que si el volcado a la hoja de Google no se puede hacer, el
archivo descargable sea Excel con la identidad visual de la Universidad —no un
CSV pelado—.

Lo que NO cambia es la regla: el sello sigue mandando. Un servicio confidencial
no exporta su historial, ni siquiera para sí mismo, porque el archivo sale de
la Unidad igual que la hoja compartida. La diferencia es que ahora hay que
decirlo servicio por servicio, en su propia pantalla.
"""

import io

import pytest
from django.core.exceptions import ValidationError
from django.test import Client
from django.urls import reverse

from apps.expediente.tests.factories import (
    crear_atencion,
    crear_estructura,
    crear_expediente,
    crear_profesional,
)
from apps.reportes import exportacion

CLAVE = "clave-larga-12345"


@pytest.fixture
def escenario(db):
    from apps.core.models import Servicio

    est = crear_estructura()
    odonto = Servicio.objects.create(
        codigo="odontologia", nombre="Odontología", seccion=est["salud"]
    )
    medico, perfil_med = crear_profesional("med_xls", est["medicina"], est["salud"])
    psicologo, perfil_psi = crear_profesional("psi_xls", est["psicologia"], est["psico"])
    # Un profesional con DOS servicios: es el caso que obliga a acotar.
    doble, perfil_doble = crear_profesional("doble_xls", est["medicina"], est["salud"])
    perfil_doble.servicios.add(odonto)

    for usuario in (medico, psicologo, doble):
        usuario.set_password(CLAVE)
        usuario.save()

    expediente = crear_expediente()
    crear_atencion(expediente, est["medicina"], perfil_med)
    crear_atencion(expediente, odonto, perfil_doble)
    crear_atencion(expediente, est["psicologia"], perfil_psi)
    return {
        "est": est,
        "odonto": odonto,
        "medico": medico,
        "psicologo": psicologo,
        "doble": doble,
        "expediente": expediente,
    }


def _cliente(usuario):
    cliente = Client()
    assert cliente.login(username=usuario.username, password=CLAVE)
    return cliente


def _abrir(contenido):
    from openpyxl import load_workbook

    return load_workbook(io.BytesIO(contenido))


# ------------------------------------------------------- acotar por servicio


@pytest.mark.django_db
def test_se_exporta_el_historial_de_UN_servicio(escenario):
    """
    Quien atiende en dos servicios no quiere el revuelto de los dos: exporta el
    de la pantalla en la que está.
    """
    filas = exportacion.historial(escenario["doble"], servicio="odontologia")
    assert {f["Servicio"] for f in filas} == {"Odontología"}


@pytest.mark.django_db
def test_sin_acotar_salen_todos_los_servicios_propios(escenario):
    filas = exportacion.historial(escenario["doble"])
    assert {f["Servicio"] for f in filas} == {"Medicina", "Odontología"}


@pytest.mark.django_db
def test_no_se_exporta_el_historial_de_un_servicio_ajeno(escenario):
    """
    Acotar no puede convertirse en un modo de pedir lo que no es tuyo: el
    parámetro llega por la URL.
    """
    filas = exportacion.historial(escenario["medico"], servicio="odontologia")
    assert filas == []


@pytest.mark.django_db
def test_un_servicio_confidencial_no_exporta_su_historial(escenario):
    """
    Ni siquiera para sí mismo. El archivo sale de la Unidad igual que la hoja
    compartida, y el sello dice «no accesible fuera del servicio».
    """
    with pytest.raises(ValidationError, match="no exporta"):
        exportacion.verificar_exportable("psicologia")


# ------------------------------------------------------------- el archivo


@pytest.mark.django_db
def test_se_descarga_un_xlsx_de_verdad(escenario):
    respuesta = _cliente(escenario["medico"]).get(
        reverse("reportes:exportar_hoja"), {"formato": "xlsx"}
    )
    assert respuesta.status_code == 200
    assert "spreadsheetml" in respuesta["Content-Type"], respuesta["Content-Type"]
    libro = _abrir(respuesta.content)
    assert libro.active.max_row >= 2, "el libro salió sin filas"


@pytest.mark.django_db
def test_el_libro_lleva_la_linea_grafica_de_la_universidad(escenario):
    """
    Verde de la UNL en la cabecera y el nombre de la Universidad en el
    encabezado: si esto se afloja, el archivo deja de parecer institucional y
    quien lo recibe no sabe de dónde salió.
    """
    from apps.core.xlsx import VERDE_UNL

    respuesta = _cliente(escenario["medico"]).get(
        reverse("reportes:exportar_hoja"), {"formato": "xlsx"}
    )
    hoja = _abrir(respuesta.content).active

    texto = " ".join(str(c.value) for fila in hoja.iter_rows(max_row=6) for c in fila if c.value)
    assert "Universidad Nacional de Loja" in texto
    assert "Bienestar Universitario" in texto

    cabecera = _fila_de_encabezados(hoja)
    assert cabecera is not None, "no se encontró la fila de encabezados"
    primera = hoja.cell(row=cabecera, column=1)
    assert VERDE_UNL.upper() in str(primera.fill.start_color.rgb).upper()
    assert primera.font.bold


def _fila_de_encabezados(hoja):
    for fila in range(1, 12):
        if hoja.cell(row=fila, column=1).value == "Fecha":
            return fila
    return None


@pytest.mark.django_db
def test_el_nombre_del_archivo_dice_de_qué_servicio_es(escenario):
    respuesta = _cliente(escenario["doble"]).get(
        reverse("reportes:exportar_hoja"), {"formato": "xlsx", "servicio": "odontologia"}
    )
    assert "odontologia" in respuesta["Content-Disposition"]


@pytest.mark.django_db
def test_el_xlsx_tampoco_lleva_psicologia(escenario):
    respuesta = _cliente(escenario["medico"]).get(
        reverse("reportes:exportar_hoja"), {"formato": "xlsx"}
    )
    hoja = _abrir(respuesta.content).active
    volcado = " ".join(
        str(c.value) for fila in hoja.iter_rows() for c in fila if c.value is not None
    )
    assert "Psicología" not in volcado


@pytest.mark.django_db
def test_la_descarga_xlsx_queda_auditada(escenario):
    from apps.auditoria.models import LogAuditoria

    LogAuditoria.objects.all().delete()
    _cliente(escenario["medico"]).get(reverse("reportes:exportar_hoja"), {"formato": "xlsx"})
    registro = LogAuditoria.objects.filter(accion=LogAuditoria.Accion.EXPORT).first()
    assert registro is not None
    assert registro.detalle["formato"] == "xlsx"


# ------------------------------------------------ el botón en cada servicio


@pytest.mark.django_db
@pytest.mark.parametrize(
    "ruta,codigo", [("medicina:bandeja", "medicina"), ("odontologia:bandeja", "odontologia")]
)
def test_cada_bandeja_ofrece_exportar_SU_historial(ruta, codigo, escenario):
    """
    El enlace tiene que ir acotado al servicio de la bandeja. Buscar solo la
    palabra «exportar» no sirve: la trae el menú lateral en todas las páginas,
    así que la prueba pasaría sin que el botón existiera.
    """
    contenido = _cliente(escenario["doble"]).get(reverse(ruta)).content.decode()
    assert f"servicio={codigo}" in contenido, f"{ruta} no ofrece exportar lo suyo"


@pytest.mark.django_db
def test_la_bandeja_de_psicologia_no_la_ofrece(escenario):
    """
    Ofrecer un botón que va a decir que no se puede es peor que no ofrecerlo:
    hace que el profesional lo intente y se lleve el aviso cada vez.
    """
    contenido = _cliente(escenario["psicologo"]).get(reverse("psicologia:bandeja")).content
    assert "exportar historial" not in contenido.decode().lower()
